#!/usr/bin/env python3
"""Convert the confirmed family timetable workbook into the app's weekly template."""
import json
from pathlib import Path
from openpyxl import load_workbook

SOURCE = Path('/Users/mb_air_2023/.hermes/cache/family-timetable/현이_삼형제_2학기_확정_주간스케쥴표_A2.xlsx')
OUTPUT = Path(__file__).resolve().parents[1] / 'weekly-template.json'
STUDENTS = [('옥승현', 'hyeon1'), ('옥수현', 'hyeon2'), ('옥서현', 'hyeon3')]

HOME = ('기상', '샤워', '식사', '등교', '귀가')
REST = ('휴식', '자유시간', '취침')
READING = ('독서',)
ARTS = ('피아노', '첼로', '플룻', '오케스트라', '음악', '태권도', '골프', '체육', '스포츠', '미술')


def category(label, weekday, slot):
    if label.upper() == 'TEST':
        return 'test'
    if any(word in label for word in HOME):
        return 'home'
    if any(word in label for word in REST):
        return 'rest'
    if any(word in label for word in READING):
        return 'reading'
    # Weekday 09:00–16:00 entries are curriculum blocks, not self-directed study records.
    if weekday < 5 and 4 <= slot < 18:
        return 'school'
    if any(word in label for word in ARTS):
        return 'arts'
    return 'study'


def slot_time(slot):
    minutes = 7 * 60 + slot * 30
    return f'{(minutes // 60) % 24:02d}:{minutes % 60:02d}'


def main():
    workbook = load_workbook(SOURCE, data_only=False)
    sheet = workbook['한눈 시간표']
    spans = {(merged.min_row, merged.min_col): merged.max_row - merged.min_row + 1 for merged in sheet.merged_cells.ranges if merged.min_col >= 2 and merged.min_row >= 5}
    events = []
    for weekday in range(7):
        for person_index, (_, student_id) in enumerate(STUDENTS):
            column = 2 + weekday * 3 + person_index
            row = 5
            while row <= 37:
                label = sheet.cell(row, column).value
                span = spans.get((row, column), 1)
                if label not in (None, '', '-'):
                    label = str(label).strip()
                    start_slot = row - 5
                    end_slot = min(33, start_slot + span)
                    events.append({
                        'id': f'{student_id}-d{weekday}-s{start_slot}',
                        'studentId': student_id,
                        'weekday': weekday,
                        'start': slot_time(start_slot),
                        'end': slot_time(end_slot) if end_slot < 33 else '23:30',
                        'title': label,
                        'category': category(label, weekday, start_slot),
                    })
                row += span
    OUTPUT.write_text(json.dumps(events, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'output': str(OUTPUT), 'events': len(events)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
